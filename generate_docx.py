import json
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX

def process_item(item, doc, level=1):
    name = item.get('name', 'Unnamed')
    
    # Check if it's a folder (has 'item' key)
    if 'item' in item:
        # Cap heading level at 9 to avoid ValueError in python-docx
        h_level = min(level, 9)
        doc.add_heading(name, level=h_level)
        for child in item['item']:
            process_item(child, doc, level + 1)
    # Otherwise it's a request (has 'request' key)
    elif 'request' in item:
        h_level = min(level, 9)
        doc.add_heading(name, level=h_level)
        req = item['request']
        
        # Method and URL
        method = req.get('method', 'GET')
        url = ""
        if isinstance(req.get('url'), dict):
            url = req['url'].get('raw', '')
        elif isinstance(req.get('url'), str):
            url = req.get('url')
            
        p = doc.add_paragraph()
        run = p.add_run(f"[{method}] ")
        run.bold = True
        if method == "GET":
            run.font.color.rgb = RGBColor(0, 100, 0)
        elif method == "POST":
            run.font.color.rgb = RGBColor(200, 150, 0)
        elif method == "PUT":
            run.font.color.rgb = RGBColor(0, 0, 200)
        elif method == "DELETE":
            run.font.color.rgb = RGBColor(200, 0, 0)
            
        p.add_run(url)
        
        # Headers
        headers = req.get('header', [])
        if headers:
            doc.add_paragraph("Headers:", style="List Bullet")
            for h in headers:
                if isinstance(h, dict):
                    k = h.get('key', '')
                    v = h.get('value', '')
                    doc.add_paragraph(f"  {k}: {v}")
                
        # Body
        body = req.get('body', {})
        if body:
            mode = body.get('mode')
            if mode == 'raw':
                raw_body = body.get('raw', '')
                if raw_body:
                    doc.add_paragraph("Body (raw):", style="List Bullet")
                    # Use a basic font for code
                    bp = doc.add_paragraph(raw_body)
                    for run in bp.runs:
                        run.font.name = 'Courier New'
            elif mode == 'formdata':
                formdata = body.get('formdata', [])
                if formdata:
                    doc.add_paragraph("Body (form-data):", style="List Bullet")
                    for fd in formdata:
                        k = fd.get('key', '')
                        v = fd.get('value', '')
                        doc.add_paragraph(f"  {k}: {v}")
            elif mode == 'urlencoded':
                urlencoded = body.get('urlencoded', [])
                if urlencoded:
                    doc.add_paragraph("Body (x-www-form-urlencoded):", style="List Bullet")
                    for fd in urlencoded:
                        k = fd.get('key', '')
                        v = fd.get('value', '')
                        doc.add_paragraph(f"  {k}: {v}")
                        
        doc.add_paragraph("-" * 40)

def main():
    with open('CAB-Booking-System-12-Levels.postman_collection.json', 'r', encoding='utf-8') as f:
        data = json.load(f)
        
    doc = Document()
    doc.add_heading(data.get('info', {}).get('name', 'Postman Collection'), 0)
    
    for item in data.get('item', []):
        process_item(item, doc, level=1)
        
    doc.save('CAB-Booking-System-Test-Cases.docx')
    print("Successfully created CAB-Booking-System-Test-Cases.docx")

if __name__ == '__main__':
    main()
